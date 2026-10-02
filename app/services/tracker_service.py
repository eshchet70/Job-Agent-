"""
Tracker Service — handles loading, querying, and updating the Master Job Application Tracker.
Reads from and writes to Final_Job_Application_Tracker-5.xlsx (Applications sheet)
with full fallback support and SQLite cross-referencing.
"""
from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd

from app.config import settings

logger = logging.getLogger(__name__)

# Expected columns in Final_Job_Application_Tracker-5.xlsx (Applications sheet)
EXPECTED_COLUMNS = [
    "Date Applied",
    "Company",
    "Role Title",
    "Level",
    "Country",
    "City / Remote",
    "Fit Score (%)",
    "Channel",
    "Referral / Contact Name",
    "Sponsorship Question?",
    "Status",
    "Response Date",
    "Days to Response",
    "Next Action",
    "Next Action Date",
    "Notes",
    "Interview/Screen Date(s)",
    "Gap Category",
    "Key Requirements",
]


def get_tracker_file_path() -> Path:
    """Resolve the active Excel tracker spreadsheet path."""
    candidates = [
        Path("Job Tracker/Final_Job_Application_Tracker-5.xlsx"),
        Path("/Users/eshchet/DEV/Job Agent/job-search-agent/Job Tracker/Final_Job_Application_Tracker-5.xlsx"),
        Path("/Users/eshchet/Documents/Job Search Tracker/Final_Job_Application_Tracker-5.xlsx"),
        Path("data/tracker_template.xlsx"),
        Path(settings.tracker_path),
    ]
    for c in candidates:
        if c.exists():
            return c.resolve()

    # Search for any tracker xlsx in workspace
    for match in Path(".").glob("**/Final_Job_Application_Tracker*.xlsx"):
        if match.is_file():
            return match.resolve()

    for match in Path(".").glob("**/*Tracker*.xlsx"):
        if match.is_file() and not match.name.startswith("~$"):
            return match.resolve()

    return Path("Job Tracker/Final_Job_Application_Tracker-5.xlsx").resolve()


def load_applications_data(
    file_path: Optional[Path | str] = None,
    include_sqlite_links: bool = True,
) -> pd.DataFrame:
    """
    Load the Applications sheet from the Excel tracker into a cleaned pandas DataFrame.
    Includes row index mappings (_excel_row) and SQLite job cross-references.
    """
    path = Path(file_path) if file_path else get_tracker_file_path()
    if not path.exists():
        logger.warning(f"Tracker file not found at {path}, returning empty DataFrame.")
        return pd.DataFrame(columns=EXPECTED_COLUMNS + ["_excel_row", "sqlite_job_id"])

    # Load via pandas
    df = pd.read_excel(path, sheet_name="Applications")

    # Clean whitespace in column names
    df.columns = [str(c).strip() for c in df.columns]

    # Ensure all expected columns exist
    for col in EXPECTED_COLUMNS:
        if col not in df.columns:
            df[col] = None

    # Track 1-based Excel row number (row 1 is header, data starts at row 2)
    df["_excel_row"] = df.index + 2

    # Parse Dates cleanly
    df["Date Applied Parsed"] = pd.to_datetime(df["Date Applied"], errors="coerce")
    df["Date Applied Clean"] = df["Date Applied Parsed"].dt.date

    df["Response Date Parsed"] = pd.to_datetime(df["Response Date"], errors="coerce")
    df["Response Date Clean"] = df["Response Date Parsed"].dt.date

    df["Next Action Date Parsed"] = pd.to_datetime(df["Next Action Date"], errors="coerce")
    df["Next Action Date Clean"] = df["Next Action Date Parsed"].dt.date

    # Clean numeric fields
    df["Fit Score (%)"] = pd.to_numeric(df["Fit Score (%)"], errors="coerce")
    df["Days to Response"] = pd.to_numeric(df["Days to Response"], errors="coerce")

    # Clean string fields
    for text_col in ["Company", "Role Title", "Level", "Country", "City / Remote", "Channel", "Status", "Notes", "Next Action"]:
        df[text_col] = df[text_col].fillna("").astype(str).str.strip()

    # Link with SQLite jobs if requested
    df["sqlite_job_id"] = None
    if include_sqlite_links:
        try:
            from app.db.repository import list_jobs
            sqlite_jobs = list_jobs(limit=500)
            if sqlite_jobs:
                sqlite_lookup = {}
                for j in sqlite_jobs:
                    c_clean = (j.company or "").strip().lower()
                    t_clean = (j.title or "").strip().lower()
                    if c_clean:
                        sqlite_lookup[(c_clean, t_clean)] = j.id
                        if c_clean not in sqlite_lookup:
                            sqlite_lookup[c_clean] = j.id

                for idx, row in df.iterrows():
                    c_val = str(row["Company"]).lower().strip()
                    t_val = str(row["Role Title"]).lower().strip()
                    if (c_val, t_val) in sqlite_lookup:
                        df.at[idx, "sqlite_job_id"] = sqlite_lookup[(c_val, t_val)]
                    elif c_val in sqlite_lookup:
                        df.at[idx, "sqlite_job_id"] = sqlite_lookup[c_val]
        except Exception as exc:
            logger.debug(f"Could not cross-reference SQLite jobs: {exc}")

    # Default sort by Date Applied descending (newest submissions first)
    df = df.sort_values(by="Date Applied Parsed", ascending=False, na_position="last").reset_index(drop=True)
    return df


def get_tracker_kpis(df: pd.DataFrame) -> dict[str, Any]:
    """Calculate key performance indicators and aggregate metrics from tracker data."""
    if df.empty:
        return {
            "total": 0,
            "applied": 0,
            "screening": 0,
            "interviewing": 0,
            "rejected": 0,
            "ghosted": 0,
            "closed": 0,
            "avg_response_days": 0.0,
            "median_response_days": 0.0,
            "last_30_days": 0,
            "last_7_days": 0,
            "max_date": None,
            "min_date": None,
        }

    total = len(df)
    status_series = df["Status"].str.lower()

    applied = int(status_series.str.contains("applied").sum())
    screening = int(status_series.str.contains("screening").sum())
    interviewing = int(status_series.str.contains("interview").sum())
    rejected = int(status_series.str.contains("reject").sum())
    ghosted = int(status_series.str.contains("ghost").sum())
    closed = int(status_series.str.contains("closed").sum())

    # Response days statistics
    valid_resp = df["Days to Response"].dropna()
    avg_resp = float(valid_resp.mean()) if not valid_resp.empty else 0.0
    med_resp = float(valid_resp.median()) if not valid_resp.empty else 0.0

    # Submissions in recent time windows
    parsed_dates = df["Date Applied Parsed"].dropna()
    max_d = parsed_dates.max() if not parsed_dates.empty else None
    min_d = parsed_dates.min() if not parsed_dates.empty else None

    # Reference date: today or latest application date in dataset
    ref_date = pd.Timestamp.now()
    if max_d is not None and max_d > ref_date:
        ref_date = max_d

    last_30 = int((parsed_dates >= (ref_date - pd.Timedelta(days=30))).sum()) if not parsed_dates.empty else 0
    last_7 = int((parsed_dates >= (ref_date - pd.Timedelta(days=7))).sum()) if not parsed_dates.empty else 0

    return {
        "total": total,
        "applied": applied,
        "screening": screening,
        "interviewing": interviewing,
        "rejected": rejected,
        "ghosted": ghosted,
        "closed": closed,
        "active_pipeline": applied + screening + interviewing,
        "avg_response_days": round(avg_resp, 1),
        "median_response_days": round(med_resp, 1),
        "last_30_days": last_30,
        "last_7_days": last_7,
        "max_date": max_d.strftime("%Y-%m-%d") if max_d is not None else None,
        "min_date": min_d.strftime("%Y-%m-%d") if min_d is not None else None,
    }


def update_application_record(
    excel_row: int,
    status: str,
    response_date: Optional[date | datetime | str] = None,
    next_action: Optional[str] = None,
    notes: Optional[str] = None,
    file_path: Optional[Path | str] = None,
) -> bool:
    """
    Update an existing application row in the Excel tracker.
    Preserves all other sheets, styles, formulas, and formatting.
    """
    path = Path(file_path) if file_path else get_tracker_file_path()
    if not path.exists():
        raise FileNotFoundError(f"Tracker file not found at {path}")

    wb = openpyxl.load_workbook(path)
    if "Applications" not in wb.sheetnames:
        raise ValueError(f"'Applications' sheet missing in {path}")

    ws = wb["Applications"]
    header_map = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}

    # Verify target row
    if excel_row < 2 or excel_row > ws.max_row:
        raise ValueError(f"Invalid row number: {excel_row} (max is {ws.max_row})")

    # Update Status
    if "Status" in header_map and status:
        ws.cell(row=excel_row, column=header_map["Status"], value=status)

    # Update Response Date
    if "Response Date" in header_map and response_date is not None:
        if isinstance(response_date, (datetime, date)):
            resp_val = datetime.combine(response_date, datetime.min.time()) if isinstance(response_date, date) else response_date
        elif isinstance(response_date, str) and response_date.strip():
            try:
                resp_val = datetime.strptime(response_date.strip(), "%Y-%m-%d")
            except Exception:
                resp_val = response_date.strip()
        else:
            resp_val = None
        ws.cell(row=excel_row, column=header_map["Response Date"], value=resp_val)

        # Recompute Days to Response if Date Applied exists
        if "Date Applied" in header_map and "Days to Response" in header_map and isinstance(resp_val, (datetime, date)):
            applied_cell_val = ws.cell(row=excel_row, column=header_map["Date Applied"]).value
            if isinstance(applied_cell_val, (datetime, date)):
                delta_days = (resp_val.date() if isinstance(resp_val, datetime) else resp_val) - (
                    applied_cell_val.date() if isinstance(applied_cell_val, datetime) else applied_cell_val
                )
                ws.cell(row=excel_row, column=header_map["Days to Response"], value=max(0, delta_days.days))

    # Update Next Action
    if "Next Action" in header_map and next_action is not None:
        ws.cell(row=excel_row, column=header_map["Next Action"], value=next_action.strip())

    # Update Notes
    if "Notes" in header_map and notes is not None:
        ws.cell(row=excel_row, column=header_map["Notes"], value=notes.strip())

    wb.save(path)
    return True


def append_application_record(
    record: dict[str, Any],
    file_path: Optional[Path | str] = None,
) -> int:
    """
    Append a new application record row to the Applications sheet.
    Returns the newly created Excel row number.
    """
    path = Path(file_path) if file_path else get_tracker_file_path()
    if not path.exists():
        raise FileNotFoundError(f"Tracker file not found at {path}")

    wb = openpyxl.load_workbook(path)
    if "Applications" not in wb.sheetnames:
        raise ValueError(f"'Applications' sheet missing in {path}")

    ws = wb["Applications"]
    header_map = {str(cell.value).strip(): cell.column for cell in ws[1] if cell.value}
    next_row = ws.max_row + 1

    for col_name, value in record.items():
        if col_name in header_map:
            col_idx = header_map[col_name]
            # Convert date objects to datetime for Excel
            if isinstance(value, date) and not isinstance(value, datetime):
                value = datetime.combine(value, datetime.min.time())
            ws.cell(row=next_row, column=col_idx, value=value)

    wb.save(path)
    return next_row
