"""
Document handler service for Word (.docx) and PDF (.pdf) parsing and generation.
"""
from __future__ import annotations

import io
from typing import Optional


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract clean text from uploaded files: .docx, .pdf, .txt, .md."""
    fn_lower = filename.lower()

    if fn_lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    elif fn_lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    else:
        # Default text/markdown decoding
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                return file_bytes.decode(enc)
            except UnicodeDecodeError:
                continue
        return file_bytes.decode("utf-8", errors="replace")


def _extract_docx(file_bytes: bytes) -> str:
    """Extract complete text from Word .docx file including paragraphs, tables, and text boxes."""
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    sections_text: list[str] = []

    # 1. Headers (often contains candidate name, contact, LinkedIn)
    for section in doc.sections:
        for hp in section.header.paragraphs:
            htxt = hp.text.strip()
            if htxt and htxt not in sections_text:
                sections_text.append(htxt)

    # 2. Main body paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt and txt not in sections_text:
            sections_text.append(txt)

    # 3. Table cells (multi-column resume layouts frequently use tables)
    for table in doc.tables:
        for row in table.rows:
            row_parts = []
            for cell in row.cells:
                cell_paras = [cp.text.strip() for cp in cell.paragraphs if cp.text.strip()]
                if cell_paras:
                    row_parts.append("\n".join(cell_paras))
                elif cell.text.strip():
                    row_parts.append(cell.text.strip())
            # Deduplicate horizontally adjacent identical cell references from merged cells
            deduped_row = []
            for item in row_parts:
                if not deduped_row or item != deduped_row[-1]:
                    deduped_row.append(item)
            for cell_item in deduped_row:
                if cell_item and cell_item not in sections_text:
                    sections_text.append(cell_item)

    # 4. Drawing Text Boxes (w:txbxContent) - essential for template resumes
    try:
        for txbx in doc.element.xpath(".//w:txbxContent"):
            for p in txbx.xpath(".//w:p"):
                runs_text = "".join(t.text for t in p.xpath(".//w:t") if t.text).strip()
                if runs_text and runs_text not in sections_text:
                    sections_text.append(runs_text)
    except Exception:
        pass

    return "\n\n".join(sections_text)


def _extract_pdf(file_bytes: bytes) -> str:
    """Extract complete text from PDF using pdfplumber with pypdf fallback."""
    pages_text: list[str] = []

    # Primary: pdfplumber (superior for multi-column layouts, tables, and visual blocks)
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                txt = page.extract_text(layout=True)
                if not txt or len(txt.strip()) < 20:
                    txt = page.extract_text()
                if txt and txt.strip():
                    pages_text.append(txt.strip())
    except Exception:
        pass

    # Fallback: pypdf if pdfplumber is empty or missed content
    if not pages_text or sum(len(p.split()) for p in pages_text) < 50:
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pypdf_pages = []
            for page in reader.pages:
                txt = ""
                try:
                    txt = page.extract_text(extraction_mode="layout") or ""
                except Exception:
                    pass
                if not txt.strip():
                    txt = page.extract_text() or ""
                if txt.strip():
                    pypdf_pages.append(txt.strip())
            if pypdf_pages and sum(len(p.split()) for p in pypdf_pages) > sum(len(p.split()) for p in pages_text):
                pages_text = pypdf_pages
        except Exception:
            pass

    return "\n\n".join(pages_text)


def create_docx(text: str, title: str = "Resume") -> bytes:
    """Generate a clean, professional Word (.docx) document from text or markdown."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document()

    # Set 0.75 in margins for resume formatting
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    primary_color = RGBColor(30, 58, 95)  # Navy
    text_color = RGBColor(30, 41, 59)     # Slate 800

    lines = text.split("\n")
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = primary_color
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)

        elif line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[3:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = primary_color
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)

        elif line.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(line[4:].strip())
            run.font.size = Pt(11)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

        elif line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            content = line[2:].strip()
            # Strip bold asterisks for clean run
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            _add_formatted_runs(p, content, text_color)

        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _add_formatted_runs(p, line, text_color)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _add_formatted_runs(paragraph, text: str, default_color):
    """Helper to handle basic **bold** segments in paragraph."""
    import re
    from docx.shared import Pt

    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.font.bold = True
        else:
            run = paragraph.add_run(part)
        run.font.size = Pt(10)
        run.font.color.rgb = default_color


def create_pdf(text: str, title: str = "Resume") -> bytes:
    """Generate a clean, styled PDF document from text or markdown."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    import re

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Normal"],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#1e3a5f"),
        fontName="Helvetica-Bold",
        spaceAfter=4,
    )
    h2_style = ParagraphStyle(
        "ResumeH2",
        parent=styles["Normal"],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e3a5f"),
        fontName="Helvetica-Bold",
        spaceBefore=8,
        spaceAfter=3,
    )
    h3_style = ParagraphStyle(
        "ResumeH3",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=5,
        spaceAfter=2,
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=styles["Normal"],
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=2,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=styles["Normal"],
        fontSize=9.5,
        leading=13.5,
        leftIndent=15,
        firstLineIndent=-10,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=1.5,
    )

    story = []
    lines = text.split("\n")

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        # XML escape
        safe_line = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        # Markdown bold conversion
        safe_line = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", safe_line)

        if line.startswith("# "):
            content = safe_line[2:].strip()
            story.append(Paragraph(content, title_style))
        elif line.startswith("## "):
            content = safe_line[3:].strip()
            story.append(Paragraph(content, h2_style))
        elif line.startswith("### "):
            content = safe_line[4:].strip()
            story.append(Paragraph(content, h3_style))
        elif line.startswith("- ") or line.startswith("* ") or line.startswith("• "):
            content = "• " + safe_line[2:].strip()
            story.append(Paragraph(content, bullet_style))
        else:
            story.append(Paragraph(safe_line, body_style))

    doc.build(story)
    return out.getvalue()
